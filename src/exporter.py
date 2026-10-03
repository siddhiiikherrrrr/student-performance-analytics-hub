"""Write analysis outputs: a styled Excel workbook plus flat CSVs for Power BI."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _style_sheet(ws):
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="22D3EE")
        c.fill = PatternFill("solid", fgColor="0A1128")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
    for i, col in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 10), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_excel(sheets: dict[str, pd.DataFrame], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, frame in sheets.items():
            frame.to_excel(xw, sheet_name=name[:31], index=False)
        for ws in xw.book.worksheets:
            _style_sheet(ws)


def write_csvs(tables: dict[str, pd.DataFrame], out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, frame in tables.items():
        frame.to_csv(out_dir / f"{name}.csv", index=False, encoding="utf-8-sig")
