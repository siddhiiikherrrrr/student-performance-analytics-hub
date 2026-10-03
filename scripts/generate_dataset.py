"""Generate a realistic 500-row student dataset (with deliberate data-quality issues).

The data is synthetic. Scores are driven by attendance, study hours and a latent
ability term so the analytics produce believable relationships, and messy rows
are injected so the Data Cleaning Module has genuine work to do.

Run:  python scripts/generate_dataset.py
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

SEED = 42
N_UNIQUE = 492          # + 8 injected duplicates = 500 raw rows
OUT = Path(__file__).resolve().parents[1] / "data" / "raw" / "student_performance_raw.xlsx"

MALE = ["Aarav", "Vihaan", "Arjun", "Reyansh", "Ishaan", "Kabir", "Rohan", "Aditya", "Dev", "Karan",
        "Yash", "Nikhil", "Rahul", "Siddharth", "Manav", "Harsh", "Krish", "Pranav", "Dhruv", "Tanmay",
        "Ayaan", "Om", "Veer", "Neel", "Jay", "Raj", "Parth", "Mihir", "Kunal", "Sahil"]
FEMALE = ["Aanya", "Diya", "Ananya", "Isha", "Meera", "Priya", "Riya", "Saanvi", "Kavya", "Nisha",
          "Tara", "Zara", "Anika", "Sneha", "Pooja", "Naina", "Aditi", "Ira", "Myra", "Kiara",
          "Avni", "Shreya", "Trisha", "Vanya", "Ritika", "Sara", "Pihu", "Jiya", "Bhavya", "Navya"]
LAST = ["Sharma", "Patel", "Mehta", "Desai", "Shah", "Joshi", "Verma", "Gupta", "Nair", "Iyer",
        "Reddy", "Kapoor", "Malhotra", "Bhatt", "Trivedi", "Chauhan", "Singh", "Kulkarni", "Pandey", "Jain",
        "Agarwal", "Thakur", "Modi", "Parekh", "Vora", "Rana", "Sinha", "Menon", "Bose", "Das"]

ARCH = {  # attendance mean/sd, study-hours mean/sd
    "high":       (93, 5, 6.0, 0.9),
    "mid":        (85, 6, 4.0, 1.0),
    "low":        (76, 7, 2.8, 0.9),
    "struggling": (64, 9, 1.6, 0.7),
}


def build() -> pd.DataFrame:
    rng = np.random.default_rng(SEED)
    n = N_UNIQUE
    gender = rng.choice(["Male", "Female"], n)
    arch = rng.choice(list(ARCH), n, p=[0.18, 0.50, 0.20, 0.12])

    att = np.array([rng.normal(ARCH[a][0], ARCH[a][1]) for a in arch]).clip(35, 100).round(1)
    hrs = np.array([rng.normal(ARCH[a][2], ARCH[a][3]) for a in arch]).clip(0.5, 10).round(1)
    ability = rng.normal(0, 7, n)
    base = 24 + 0.30 * att + 4.2 * hrs + ability
    is_m = (gender == "Male")

    math = base - 1 + 2.0 * is_m + rng.normal(0, 6, n)
    sci = base - 1 + 0.5 * is_m + rng.normal(0, 6, n)
    eng = base + 2 - 3.0 * is_m + rng.normal(0, 6, n)

    names, used = [], set()
    while len(names) < n:
        i = len(names)
        first = rng.choice(MALE if gender[i] == "Male" else FEMALE)
        full = f"{first} {rng.choice(LAST)}"
        if full not in used:
            used.add(full)
            names.append(full)

    df = pd.DataFrame({
        "Student_ID": [f"STU-{i:04d}" for i in range(1, n + 1)],
        "Student_Name": names, "Gender": gender,
        "Attendance_Percentage": att, "Study_Hours": hrs,
        "Math_Score": math.clip(8, 100).round(1),
        "Science_Score": sci.clip(8, 100).round(1),
        "English_Score": eng.clip(8, 100).round(1),
    })

    # ---- inject realistic mess -------------------------------------------
    def pick(k):
        return rng.choice(n, k, replace=False)

    for col, k in [("Attendance_Percentage", 12), ("Study_Hours", 10), ("Math_Score", 8),
                   ("Science_Score", 6), ("English_Score", 6)]:
        df.loc[pick(k), col] = np.nan
    df.loc[pick(2), "Attendance_Percentage"] = [108.0, 115.0]     # impossible
    df.loc[pick(1), "Math_Score"] = 105.0
    df.loc[pick(1), "Science_Score"] = -5.0
    df.loc[pick(1), "Study_Hours"] = -2.0
    df["Gender"] = df["Gender"].astype(object)
    for idx, val in zip(pick(20), ["M", "f", "female ", " MALE", "F", "m"] * 4):
        df.at[idx, "Gender"] = val
    for idx in pick(15):
        nm = df.at[idx, "Student_Name"]
        df.at[idx, "Student_Name"] = ("  " + nm.lower() + " ") if idx % 2 else nm.upper()

    dupes = df.iloc[rng.choice(n, 8, replace=False)].copy()
    dupes["Student_Name"] = dupes["Student_Name"].astype(str) + " "
    raw = pd.concat([df, dupes], ignore_index=True)
    return raw.sample(frac=1, random_state=SEED).reset_index(drop=True)


def write(raw: pd.DataFrame) -> None:
    dictionary = pd.DataFrame({
        "Field": list(raw.columns),
        "Type": ["Text (ID)", "Text", "Category", "Number (0-100)", "Number (hours/day)",
                 "Number (0-100)", "Number (0-100)", "Number (0-100)"],
        "Description": ["Unique student identifier", "Full name", "Male / Female",
                        "Share of classes attended", "Average self-study hours per day",
                        "Mathematics exam score", "Science exam score", "English exam score"],
    })
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        raw.to_excel(xw, sheet_name="Students", index=False)
        dictionary.to_excel(xw, sheet_name="Data_Dictionary", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(name="Arial", bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="0A1128")
                c.alignment = Alignment(horizontal="center")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value or "")) for c in col) + 3
            ws.freeze_panes = "A2"


if __name__ == "__main__":
    data = build()
    write(data)
    print(f"Wrote {len(data)} rows -> {OUT}")
